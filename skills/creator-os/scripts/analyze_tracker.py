#!/usr/bin/env python3
"""Ringkas performa konten dari Tracker (xlsx/csv) jadi markdown untuk dianalisis.

Pemakaian:
  analyze_tracker.py tracker.xlsx                 # semua data
  analyze_tracker.py tracker.xlsx --month 2026-09 # satu bulan
  analyze_tracker.py data.csv --map views=reach --map save=saves
  analyze_tracker.py tracker.xlsx --add-report hasil.xlsx   # salin file + tambah tab 'Laporan Bulanan' berisi rumus

ER = (likes + komentar + share + save) / views * 100  (rumus tracker asli, basis views).
Konten yang umurnya < 24 jam (YouTube < 72 jam) dari --today ditandai belum stabil dan tidak ikut top/bottom.
"""
import argparse, shutil, sys
from datetime import date, datetime, timedelta
import pandas as pd

NEED = ["tanggal", "platform", "format", "kategori", "topik", "hook", "views",
        "likes", "komentar", "share", "save", "catatan"]
METRIC = ["views", "likes", "komentar", "share", "save"]


def load(path, sheet, maps):
    if path.lower().endswith(".csv"):
        raw = pd.read_csv(path, header=None, dtype=object)
    else:
        xl = pd.ExcelFile(path)
        name = sheet or next((s for s in xl.sheet_names if "log" in s.lower()), xl.sheet_names[0])
        raw = xl.parse(name, header=None, dtype=object)
    alias = {v.lower(): k for k, v in maps.items()}
    hdr = None
    for i in range(min(20, len(raw))):
        cells = [str(x).strip().lower() for x in raw.iloc[i].tolist()]
        if "views" in cells or alias.get("views", "views") in cells:
            hdr = i
            break
    if hdr is None:
        sys.exit("Header tidak ketemu (butuh kolom 'Views' atau --map views=<nama kolom>).")
    cols = [str(x).strip().lower() for x in raw.iloc[hdr].tolist()]
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [alias.get(c, c) for c in cols]
    df = df.loc[:, ~df.columns.duplicated()]
    df = df[[c for c in NEED if c in df.columns]]
    for m in METRIC:
        df[m] = pd.to_numeric(df.get(m), errors="coerce").fillna(0) if m in df else 0
    df = df[df["views"] > 0].copy()
    t = df.get("tanggal")
    if t is not None:
        num = pd.to_numeric(t, errors="coerce")
        df["tanggal"] = pd.to_datetime(t, errors="coerce")
        serial = num.notna() & (num > 20000)
        df.loc[serial, "tanggal"] = pd.to_datetime("1899-12-30") + pd.to_timedelta(num[serial], unit="D")
    return df


def table(df, by):
    if by not in df or df[by].isna().all():
        return ""
    g = df.groupby(by).agg(konten=("views", "size"), views=("views", "sum"),
                           eng=("eng", "sum"), save=("save", "sum"), share=("share", "sum"))
    g["avg_views"] = (g["views"] / g["konten"]).round(0).astype(int)
    g["ER%"] = (g["eng"] / g["views"] * 100).round(2)
    g["save%"] = (g["save"] / g["views"] * 100).round(2)
    g["share%"] = (g["share"] / g["views"] * 100).round(2)
    g = g.sort_values("ER%", ascending=False)[["konten", "views", "avg_views", "ER%", "save%", "share%"]]
    return f"\n### Per {by}\n" + g.to_markdown() + "\n"


def item(r):
    d = r["tanggal"].strftime("%Y-%m-%d") if pd.notna(r.get("tanggal")) else "-"
    s = (f"- {d} · {r.get('platform','')} · {r.get('format','')} · {r.get('kategori','')} — "
         f"**{r.get('topik','')}** | views {int(r['views']):,} · ER {r['er']:.2f}% · "
         f"save {r['save_r']:.2f}% · share {r['share_r']:.2f}%")
    if isinstance(r.get("hook"), str):
        s += f"\n  Hook: \"{r['hook']}\""
    if isinstance(r.get("catatan"), str):
        s += f"\n  Catatan user: {r['catatan']}"
    return s


def report(df, today):
    df["eng"] = df["likes"] + df["komentar"] + df["share"] + df["save"]
    df["er"] = df["eng"] / df["views"] * 100
    df["save_r"] = df["save"] / df["views"] * 100
    df["share_r"] = df["share"] / df["views"] * 100
    if "tanggal" in df:
        age = (pd.Timestamp(today) - df["tanggal"]).dt.days
        yt = df.get("platform", pd.Series("", index=df.index)).astype(str).str.lower().str.contains("youtube")
        df["stabil"] = ~((age < 1) | (yt & (age < 3)))
        df["hari"] = df["tanggal"].dt.dayofweek.map(dict(enumerate(["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"])))
    else:
        df["stabil"] = True
    out = ["## Ringkasan",
           f"- Konten: {len(df)} · Total views: {int(df['views'].sum()):,} · Rata-rata views: {int(df['views'].mean()):,}",
           f"- ER keseluruhan: {df['eng'].sum() / df['views'].sum() * 100:.2f}% · Rata-rata ER per konten: {df['er'].mean():.2f}%"]
    if "tanggal" in df and df["tanggal"].notna().any():
        out.append(f"- Periode: {df['tanggal'].min():%Y-%m-%d} s/d {df['tanggal'].max():%Y-%m-%d}")
    unstable = int((~df["stabil"]).sum())
    if unstable:
        out.append(f"- ⚠️ {unstable} konten belum stabil (<24 jam / YouTube <72 jam) — tidak ikut top/bottom.")
    if len(df) < 8:
        out.append("- ⚠️ Data < 8 konten: perlakukan pola sebagai indikasi awal.")
    for by in ["platform", "format", "kategori", "hari"]:
        out.append(table(df, by))
    s = df[df["stabil"]].sort_values("er", ascending=False)
    n = min(3, max(1, len(s) // 2))
    out.append("\n## Top konten (ER)\n" + "\n".join(item(r) for _, r in s.head(n).iterrows()))
    out.append("\n## Bottom konten (ER)\n" + "\n".join(item(r) for _, r in s.tail(n).iloc[::-1].iterrows()))
    return "\n".join(out)


def add_report(src, dst, sheet):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    shutil.copy(src, dst)
    wb = load_workbook(dst)
    ws = wb[sheet] if sheet else next((wb[s] for s in wb.sheetnames if "log" in s.lower()), wb.worksheets[0])
    hdr_row, col = None, {}
    for row in ws.iter_rows(min_row=1, max_row=20):
        names = {str(c.value).strip().lower(): c.column for c in row if c.value is not None}
        if "views" in names:
            hdr_row, col = row[0].row, names
            break
    if not hdr_row or "tanggal" not in col:
        sys.exit("Tab log butuh kolom Tanggal dan Views untuk --add-report.")
    q = f"'{ws.title}'!"
    rng = lambda k: f"{q}${get_column_letter(col[k])}${hdr_row + 1}:${get_column_letter(col[k])}$5000"
    months = sorted({c.value.strftime("%Y-%m") for c in ws[get_column_letter(col["tanggal"])][hdr_row:]
                     if isinstance(c.value, (datetime, date))})
    if "Laporan Bulanan" in wb.sheetnames:
        del wb["Laporan Bulanan"]
    rep = wb.create_sheet("Laporan Bulanan")
    heads = ["Bulan", "Mulai", "Selesai", "Total Konten", "Total Views", "Rata-rata Views", "Total Engagement", "ER (%)"]
    plats = ["Instagram", "TikTok", "YouTube"] if "platform" in col else []
    heads += [f"Konten {p}" for p in plats]
    rep.append(heads)
    for c in rep[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    T = rng("tanggal")
    for i, m in enumerate(months, start=2):
        y, mo = map(int, m.split("-"))
        crit = f'{T},">="&B{i},{T},"<"&C{i}'
        eng = "+".join(f"SUMIFS({rng(k)},{crit})" for k in ["likes", "komentar", "share", "save"] if k in col)
        row = [m, date(y, mo, 1), date(y + mo // 12, mo % 12 + 1, 1),
               f"=COUNTIFS({crit})", f"=SUMIFS({rng('views')},{crit})",
               f"=IFERROR(E{i}/D{i},0)", f"={eng}" if eng else 0, f"=IFERROR(G{i}/E{i}*100,0)"]
        row += [f'=COUNTIFS({crit},{rng("platform")},"{p}")' for p in plats]
        rep.append(row)
        for c in (rep.cell(i, 2), rep.cell(i, 3)):
            c.number_format = "yyyy-mm-dd"
        rep.cell(i, 5).number_format = rep.cell(i, 6).number_format = "#,##0"
        rep.cell(i, 8).number_format = "0.00"
    n = len(months) + 3
    rep.cell(n, 1, "Catatan: Mulai/Selesai = rentang bulan (Selesai tidak termasuk). ER = (likes+komentar+share+save)/views×100. Semua angka rumus dari tab log.")
    for r in rep.iter_rows():
        for c in r:
            if c.font.name != "Arial":
                c.font = Font(name="Arial", size=10, bold=c.font.bold, color=c.font.color)
    for i in range(1, len(heads) + 1):
        rep.column_dimensions[get_column_letter(i)].width = 16
    wb.save(dst)
    print(f"OK: tab 'Laporan Bulanan' ditambahkan ke {dst} ({len(months)} bulan). Jalankan recalc sebelum dibagikan.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--sheet")
    ap.add_argument("--month", help="YYYY-MM")
    ap.add_argument("--today", default=date.today().isoformat())
    ap.add_argument("--map", action="append", default=[], metavar="kolom_standar=nama_di_file")
    ap.add_argument("--add-report", metavar="OUT.xlsx")
    a = ap.parse_args()
    maps = dict(m.split("=", 1) for m in a.map)
    df = load(a.path, a.sheet, maps)
    if a.month and "tanggal" in df:
        df = df[df["tanggal"].dt.strftime("%Y-%m") == a.month]
    if df.empty:
        sys.exit("Tidak ada baris dengan views > 0 untuk periode ini.")
    print(report(df, a.today))
    if a.add_report:
        add_report(a.path, a.add_report, a.sheet)
