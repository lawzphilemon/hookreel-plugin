#!/usr/bin/env python3
"""Bangun file .xlsx kalender konten dari JSON.

Input JSON:
{
  "judul": "Kalender Konten 30 Hari — Niche Keuangan",   # opsional
  "start_date": "2026-10-19",                             # opsional (YYYY-MM-DD) -> isi kolom Tanggal
  "isi_hari_libur": true,                                 # opsional, default true: hari yang tidak ada di rows jadi 'Libur / Opsional'
  "rows": [
    {"hari": 1, "format": "Carousel Edukasi", "tujuan": "TOFU",
     "topik": "...", "hook": "...", "cta": "...",
     "fase": "Fase 1", "fokus_fase": "Awareness & Credibility"}   # fase/fokus_fase opsional
  ]
}
Pemakaian: build_calendar.py rencana.json output.xlsx
"""
import json, math, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FASE = [(1, 30, "Fase 1", "Awareness & Credibility", "DCEBFA"),
        (31, 60, "Fase 2", "Engagement & Authority", "E2F2DF"),
        (61, 10**6, "Fase 3", "Conversion & Monetization", "FCE8D6")]
FUNNEL = {"TOFU": "9CC9F0", "MOFU": "A8D8A0", "BOFU": "F5B98A"}
LIBUR = "Libur / Opsional"
GREY = "E7E7E7"
HARI = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
HEAD = ["Hari ke-", "Tanggal", "Minggu ke-", "Fase", "Fokus Fase", "Format Konten",
        "Tujuan", "Topik Konten", "Hook Ide", "CTA", "Catatan"]
WIDTH = [8, 16, 10, 9, 26, 24, 10, 42, 52, 42, 24]
F = "Arial"


def fase_of(h):
    for lo, hi, name, focus, color in FASE:
        if lo <= h <= hi:
            return name, focus, color


def tanggal(start, h):
    if not start:
        return f"Hari {h}"
    d = start + timedelta(days=h - 1)
    return f"{HARI[d.weekday()]}, {d.day} {BULAN[d.month - 1]} {d.year}"


def main(src, out):
    plan = json.load(open(src, encoding="utf-8"))
    rows = {int(r["hari"]): r for r in plan["rows"]}
    start = date.fromisoformat(plan["start_date"]) if plan.get("start_date") else None
    last = max(rows)
    days = range(1, last + 1) if plan.get("isi_hari_libur", True) else sorted(rows)

    wb = Workbook()
    guide = wb.active
    guide.title = "Cara Pakai"
    cal = wb.create_sheet("Kalender")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    title = plan.get("judul", f"Kalender Konten {last} Hari")
    cal.append([title]); cal["A1"].font = Font(name=F, bold=True, size=14)
    cal.append([])
    cal.append(HEAD)
    for c in cal[3]:
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    for h in days:
        name, focus, color = fase_of(h)
        r = rows.get(h)
        if r:
            vals = [h, tanggal(start, h), math.ceil(h / 7), r.get("fase", name),
                    r.get("fokus_fase", focus), r.get("format", ""), r.get("tujuan", ""),
                    r.get("topik", ""), r.get("hook", ""), r.get("cta", ""), r.get("catatan", "")]
        else:
            vals = [h, tanggal(start, h), math.ceil(h / 7), name, focus] + [LIBUR] * 5 + [""]
        cal.append(vals)
        row = cal.max_row
        for i, c in enumerate(cal[row], 1):
            c.font = Font(name=F, size=10, color="7F7F7F" if not r else "000000")
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if i in (1, 3, 4) else None)
            c.border = border
            c.fill = PatternFill("solid", fgColor=GREY if not r else color)
        tujuan = str(vals[6]).upper()
        if r and tujuan in FUNNEL:
            cell = cal.cell(row=row, column=7)
            cell.fill = PatternFill("solid", fgColor=FUNNEL[tujuan])
            cell.font = Font(name=F, size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="top")

    for i, w in enumerate(WIDTH, 1):
        cal.column_dimensions[cal.cell(row=3, column=i).column_letter].width = w
    cal.freeze_panes = "C4"
    cal.page_setup.orientation = "landscape"
    cal.page_setup.fitToWidth, cal.page_setup.fitToHeight = 1, 0
    cal.sheet_properties.pageSetUpPr.fitToPage = True
    cal.auto_filter.ref = f"A3:K{cal.max_row}"

    lines = [
        (title, True, 14), ("", False, 10),
        ("Cara pakai", True, 11),
        ("1. Buka tab 'Kalender'. Setiap baris = satu hari. Baris abu-abu = libur / opsional.", False, 10),
        ("2. Kolom Topik, Hook Ide, dan CTA sudah terisi — edit bebas sesuai gaya bicaramu.", False, 10),
        ("3. Kolom Catatan (paling kanan) kosong untuk kamu isi: status produksi, link, insight.", False, 10),
        ("4. Pakai filter di header untuk lihat satu format atau satu tahap funnel saja.", False, 10),
        ("5. Setelah posting, catat performanya di Tracker Performa Konten (24 jam setelah posting).", False, 10),
        ("", False, 10), ("Ritme mingguan", True, 11),
        ("Senin: Carousel Edukasi (TOFU) · Selasa: Video/Reels (TOFU) · Rabu: Carousel Personal Story (MOFU)", False, 10),
        ("Kamis: Video Tutorial/Tips (MOFU) · Jumat: Carousel/Video CTA (BOFU) · Sabtu–Minggu: libur/opsional", False, 10),
        ("", False, 10), ("Panduan warna", True, 11),
    ]
    for text, bold, size in lines:
        guide.append([text]); guide.cell(row=guide.max_row, column=1).font = Font(name=F, bold=bold, size=size)
    legend = [(c, f"{n} (Hari {lo}–{hi if hi < 10**6 else '…'}): {fo}") for lo, hi, n, fo, c in FASE] + \
             [(FUNNEL["TOFU"], "TOFU — awareness, tarik audiens baru"),
              (FUNNEL["MOFU"], "MOFU — engagement, bangun kepercayaan"),
              (FUNNEL["BOFU"], "BOFU — konversi, dorong tindakan"),
              (GREY, "Libur / opsional")]
    for color, text in legend:
        guide.append(["", text])
        guide.cell(row=guide.max_row, column=1).fill = PatternFill("solid", fgColor=color)
        guide.cell(row=guide.max_row, column=2).font = Font(name=F, size=10)
    guide.column_dimensions["A"].width = 6
    guide.column_dimensions["B"].width = 70
    for row in guide.iter_rows(min_row=1, max_row=len(lines), max_col=1):
        for c in row:
            guide.merge_cells(start_row=c.row, start_column=1, end_row=c.row, end_column=2)

    wb.active = 1
    wb.save(out)
    posts = len(rows)
    print(f"OK: {out} — {last} hari, {posts} posting, {len(list(days)) - posts} libur")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
